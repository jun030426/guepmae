# scripts/write_listing_excel_test.py
import importlib.util, os, tempfile, unittest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
spec = importlib.util.spec_from_file_location(
    "write_listing_excel", os.path.join(ROOT, "scripts", "write-listing-excel.py"))
wle = importlib.util.module_from_spec(spec)
spec.loader.exec_module(wle)

FULL = {"단지명": "래미안대치팰리스", "구/시군구": "서울특별시 강남구",
        "전용면적(㎡)": 84, "호가(원)": 2650000000}


class HeaderTest(unittest.TestCase):
    def test_required_columns_come_first(self):
        h = wle.header_for([FULL])
        self.assertEqual(h[:4], ["단지명", "구/시군구", "전용면적(㎡)", "호가(원)"])

    def test_optional_columns_appended_in_canonical_order(self):
        h = wle.header_for([{**FULL, "향": "남향", "층": "12층"}])
        self.assertEqual(h[4:], ["층", "향"])  # FIELD_ORDER 순서, 등장 순서가 아님

    def test_photo_columns_expand_to_max_count(self):
        h = wle.header_for([{**FULL, "사진": ["a"]}, {**FULL, "사진": ["a", "b", "c"]}])
        self.assertEqual(h[-3:], ["사진1", "사진2", "사진3"])

    def test_no_photo_columns_when_absent(self):
        self.assertNotIn("사진1", wle.header_for([FULL]))


class BuildWorkbookTest(unittest.TestCase):
    def test_sheet_is_named_maemul(self):
        wb = wle.build_workbook([FULL])
        self.assertEqual(wb.sheetnames, ["매물입력"])

    def test_row_values_land_under_their_header(self):
        wb = wle.build_workbook([{**FULL, "향": "남향"}])
        ws = wb["매물입력"]
        header = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        self.assertEqual(row[header.index("단지명")], "래미안대치팰리스")
        self.assertEqual(row[header.index("호가(원)")], 2650000000)
        self.assertEqual(row[header.index("향")], "남향")

    def test_photos_spread_across_columns(self):
        wb = wle.build_workbook([{**FULL, "사진": ["u1", "u2"]}])
        ws = wb["매물입력"]
        header = [c.value for c in ws[1]]
        row = [c.value for c in ws[2]]
        self.assertEqual(row[header.index("사진1")], "u1")
        self.assertEqual(row[header.index("사진2")], "u2")

    def test_rows_missing_required_fields_are_dropped(self):
        wb = wle.build_workbook([FULL, {"단지명": "이름만"}])
        self.assertEqual(wb["매물입력"].max_row, 2)  # 헤더 + 1행

    def test_blank_optional_cell_is_none(self):
        wb = wle.build_workbook([{**FULL, "향": "남향"}, FULL])
        ws = wb["매물입력"]
        header = [c.value for c in ws[1]]
        self.assertIsNone([c.value for c in ws[3]][header.index("향")])


class PrivacyGuardTest(unittest.TestCase):
    def test_agent_columns_are_rejected(self):
        for banned in ("중개사무소명", "중개사 전화번호", "중개사 등록번호"):
            with self.assertRaises(ValueError, msg=banned):
                wle.build_workbook([{**FULL, banned: "x"}])


class RoundTripTest(unittest.TestCase):
    def test_saved_file_reopens_with_expected_shape(self):
        from openpyxl import load_workbook
        wb = wle.build_workbook([FULL, {**FULL, "단지명": "두번째"}])
        with tempfile.TemporaryDirectory() as td:
            p = os.path.join(td, "out.xlsx")
            wb.save(p)
            reopened = load_workbook(p, data_only=True)
        self.assertIn("매물입력", reopened.sheetnames)
        self.assertEqual(reopened["매물입력"].max_row, 3)


if __name__ == "__main__":
    unittest.main()
